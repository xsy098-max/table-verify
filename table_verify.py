import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
import os
import sys
import threading
import copy
import json

from task_model import (
    Task, DataSource, BlockConfig, VerifyGroup,
    BLOCK_TYPES, BLOCK_DESCRIPTIONS, BLOCK_PARAMS, VERSION, AUTHOR
)
from table_loader import TableLoader, TableData
from runner import Runner, RunResult, GroupResult


class BatchRunDialog(tk.Toplevel):
    def __init__(self, parent, task_names):
        super().__init__(parent)
        self.title("批量运行")
        self.geometry("400x450")
        self.resizable(True, True)
        self.result = None
        self.transient(parent)
        self.grab_set()

        main = ttk.Frame(self, padding=15)
        main.pack(fill=tk.BOTH, expand=True)

        self.select_all_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(main, text="全选 / 全不选", variable=self.select_all_var,
                         command=self._toggle_all).pack(anchor=tk.W, pady=(0, 10))

        list_frame = ttk.Frame(main)
        list_frame.pack(fill=tk.BOTH, expand=True)

        self.canvas = tk.Canvas(list_frame, highlightthickness=0)
        sb = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.canvas.yview)
        self.inner = ttk.Frame(self.canvas)
        self.canvas.configure(yscrollcommand=sb.set)
        self.canvas.create_window((0, 0), window=self.inner, anchor=tk.NW)
        self.inner.bind('<Configure>', lambda e: self.canvas.configure(scrollregion=self.canvas.bbox('all')))
        self.canvas.bind('<Configure>', lambda e: self.canvas.itemconfig(1, width=e.width))
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sb.pack(side=tk.RIGHT, fill=tk.Y)

        self.task_vars = []
        for name in task_names:
            var = tk.BooleanVar(value=True)
            self.task_vars.append((name, var))
            ttk.Checkbutton(self.inner, text=name, variable=var).pack(anchor=tk.W, pady=2)

        self.count_label = ttk.Label(main, text=f"已选择 {len(task_names)} 个任务", foreground='gray')
        self.count_label.pack(anchor=tk.W, pady=(10, 0))
        for _, var in self.task_vars:
            var.trace_add('write', lambda *a: self._update_count())

        btn_frame = ttk.Frame(main)
        btn_frame.pack(pady=(15, 0))
        ttk.Button(btn_frame, text="运行选中", command=self._ok, width=12).pack(side=tk.LEFT, padx=10)
        ttk.Button(btn_frame, text="取消", command=self.destroy, width=10).pack(side=tk.LEFT, padx=10)

    def _toggle_all(self):
        val = self.select_all_var.get()
        for _, var in self.task_vars:
            var.set(val)

    def _update_count(self):
        count = sum(1 for _, v in self.task_vars if v.get())
        self.count_label.config(text=f"已选择 {count} 个任务")

    def _ok(self):
        selected = [name for name, var in self.task_vars if var.get()]
        if not selected:
            messagebox.showwarning("提示", "请至少选择一个任务", parent=self)
            return
        self.result = selected
        self.destroy()


class DataSourceDialog(tk.Toplevel):
    def __init__(self, parent, ds=None):
        super().__init__(parent)
        self.title("编辑数据源" if ds else "添加数据源")
        self.geometry("500x320")
        self.resizable(False, False)
        self.result = None
        self.transient(parent)
        self.grab_set()

        main = ttk.Frame(self, padding=15)
        main.pack(fill=tk.BOTH, expand=True)

        ttk.Label(main, text="数据源名称:").grid(row=0, column=0, sticky=tk.W, pady=5)
        self.name_var = tk.StringVar(value=ds.name if ds else "")
        ttk.Entry(main, textvariable=self.name_var, width=40).grid(row=0, column=1, columnspan=2, sticky=tk.EW, pady=5)

        ttk.Label(main, text="文件路径:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.path_var = tk.StringVar(value=ds.file_path if ds else "")
        path_entry = ttk.Entry(main, textvariable=self.path_var, width=32)
        path_entry.grid(row=1, column=1, sticky=tk.EW, pady=5)
        ttk.Button(main, text="浏览...", command=self._browse).grid(row=1, column=2, padx=(5, 0), pady=5)

        ttk.Label(main, text="工作表:").grid(row=2, column=0, sticky=tk.W, pady=5)
        self.sheet_var = tk.StringVar(value=ds.sheet_name if ds and ds.sheet_name else "")
        self.sheet_combo = ttk.Combobox(main, textvariable=self.sheet_var, width=37, state='readonly')
        self.sheet_combo.grid(row=2, column=1, columnspan=2, sticky=tk.EW, pady=5)

        ttk.Label(main, text="表头行号:").grid(row=3, column=0, sticky=tk.W, pady=5)
        self.header_var = tk.IntVar(value=ds.header_row if ds else 2)
        ttk.Spinbox(main, from_=1, to=100, textvariable=self.header_var, width=10).grid(row=3, column=1, sticky=tk.W, pady=5)
        ttk.Label(main, text="(字段名所在行)", foreground='gray').grid(row=3, column=1, sticky=tk.W, padx=(80, 0), pady=5)

        path_entry.bind('<FocusOut>', self._update_sheets)

        btn_frame = ttk.Frame(main)
        btn_frame.grid(row=4, column=0, columnspan=3, pady=20)
        ttk.Button(btn_frame, text="确定", command=self._ok, width=10).pack(side=tk.LEFT, padx=10)
        ttk.Button(btn_frame, text="取消", command=self.destroy, width=10).pack(side=tk.LEFT, padx=10)

        main.columnconfigure(1, weight=1)
        if ds and ds.file_path:
            self._update_sheets()

    def _browse(self):
        ftypes = [("所有支持", "*.xlsx;*.xlsm;*.csv"), ("Excel", "*.xlsx;*.xlsm"), ("CSV", "*.csv"), ("所有", "*.*")]
        path = filedialog.askopenfilename(filetypes=ftypes)
        if path:
            self.path_var.set(path)
            self._update_sheets()

    def _update_sheets(self, event=None):
        path = self.path_var.get().strip()
        if not path:
            return
        ext = os.path.splitext(path)[1].lower()
        if ext in ('.xlsx', '.xlsm'):
            try:
                names = TableLoader.get_sheet_names(path)
                self.sheet_combo['values'] = names
                if names:
                    self.sheet_combo.set(names[0])
            except Exception:
                self.sheet_combo['values'] = []
        else:
            self.sheet_combo['values'] = []
            self.sheet_combo.set('')

    def _ok(self):
        name = self.name_var.get().strip()
        path = self.path_var.get().strip()
        if not name:
            messagebox.showwarning("提示", "请输入数据源名称", parent=self)
            return
        if not path:
            messagebox.showwarning("提示", "请选择文件", parent=self)
            return
        sheet = self.sheet_var.get().strip() or None
        self.result = DataSource(name=name, file_path=path, sheet_name=sheet, header_row=self.header_var.get())
        self.destroy()


class DataPreviewDialog(tk.Toplevel):
    def __init__(self, parent, ds, table_cache):
        super().__init__(parent)
        self.title(f"预览: {ds.name}")
        self.geometry("800x500")
        self.transient(parent)

        frame = ttk.Frame(self, padding=5)
        frame.pack(fill=tk.BOTH, expand=True)

        info_text = f"文件: {os.path.basename(ds.file_path)}  工作表: {ds.sheet_name or '(CSV)'}  表头行: {ds.header_row}"
        ttk.Label(frame, text=info_text, foreground='gray').pack(anchor=tk.W, pady=(0, 5))

        table = table_cache.get(ds.name)
        if table is None:
            ttk.Label(frame, text="无法加载此数据源", foreground='red').pack()
            return

        cols = table.headers[:20]
        tree = ttk.Treeview(frame, columns=cols, show='headings', height=20)
        for h in cols:
            tree.heading(h, text=h)
            tree.column(h, width=80, minwidth=40, stretch=True)

        for row_idx in range(table.header_row, min(table.header_row + 12, len(table.raw_data))):
            row_data = table.raw_data[row_idx]
            values = []
            for ci, h in enumerate(cols):
                val = row_data[ci] if ci < len(row_data) else ''
                values.append(str(val)[:50] if val else '')
            tree.insert('', tk.END, values=values)

        hsb = ttk.Scrollbar(frame, orient=tk.HORIZONTAL, command=tree.xview)
        vsb = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=tree.yview)
        tree.configure(xscrollcommand=hsb.set, yscrollcommand=vsb.set)

        hsb.pack(side=tk.BOTTOM, fill=tk.X)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)


class StepWidget(ttk.LabelFrame):
    def __init__(self, parent, block_config, app, step_index=0, **kwargs):
        super().__init__(parent, text=f"{BLOCK_TYPES.get(block_config.block_type, block_config.block_type)}", **kwargs)
        self.block_config = block_config
        self.app = app
        self.step_index = step_index
        self.param_widgets = {}
        self._row_frames = {}
        self._inner = None
        self._build_fields()

    def _get_ds_names(self):
        return [ds.name for ds in self.app.current_task.data_sources]

    def _get_var_names(self):
        group = self.app._get_current_group()
        if not group:
            return []
        names = []
        for i, step in enumerate(group.steps):
            if i >= self.step_index:
                break
            out_var = step.params.get('output_var', '')
            if out_var:
                names.append(out_var)
        return names

    def _get_columns_for_ds(self, ds_name):
        table = self.app.table_cache.get(ds_name)
        if table:
            return list(table.headers)
        return []

    def _build_fields(self):
        if self._inner:
            self._inner.destroy()
        self.param_widgets = {}
        self._row_frames = {}

        self._inner = ttk.Frame(self, padding=5)
        self._inner.pack(fill=tk.BOTH, expand=True)

        param_defs = BLOCK_PARAMS.get(self.block_config.block_type, [])
        current_mode = self.block_config.params.get('mode', '')
        if not current_mode:
            for p in param_defs:
                if p['key'] == 'mode' and p['type'] == 'choice':
                    current_mode = p.get('choices', [''])[0]
                    self.block_config.params['mode'] = current_mode
                    break
        row = 0

        for pdef in param_defs:
            key = pdef['key']
            label = pdef['label']
            ptype = pdef['type']
            condition = pdef.get('condition', None)

            if condition:
                cond_key, cond_val = condition.split('==')
                cond_key = cond_key.strip()
                cond_val = cond_val.strip()
                check_val = current_mode if cond_key == 'mode' else self.block_config.params.get(cond_key, '')
                if str(check_val) != cond_val:
                    continue

            frame = ttk.Frame(self._inner)
            frame.grid(row=row, column=0, columnspan=2, sticky=tk.EW, pady=2)
            self._row_frames[key] = frame
            ttk.Label(frame, text=f"{label}:").pack(side=tk.LEFT, padx=(0, 5))

            if ptype == 'datasource':
                var = tk.StringVar(value=self.block_config.params.get(key, ''))
                w = ttk.Combobox(frame, textvariable=var, values=self._get_ds_names(), width=30)
                w.pack(side=tk.LEFT, fill=tk.X, expand=True)
                w.bind('<<ComboboxSelected>>', lambda e, k=key: self._on_datasource_change(k))
                self.param_widgets[key] = ('combobox', var, w)

            elif ptype == 'variable':
                var = tk.StringVar(value=self.block_config.params.get(key, ''))
                w = ttk.Combobox(frame, textvariable=var, values=self._get_var_names(), width=30)
                w.pack(side=tk.LEFT, fill=tk.X, expand=True)
                self.param_widgets[key] = ('combobox', var, w)

            elif ptype == 'column':
                var = tk.StringVar(value=self.block_config.params.get(key, ''))
                ds_name = self._find_related_datasource(key)
                cols = self._get_columns_for_ds(ds_name)
                w = ttk.Combobox(frame, textvariable=var, values=cols, width=30)
                w.pack(side=tk.LEFT, fill=tk.X, expand=True)
                self.param_widgets[key] = ('combobox', var, w)

            elif ptype == 'choice':
                choices = pdef.get('choices', [])
                choice_labels = pdef.get('choice_labels', choices)
                allow_custom = pdef.get('allow_custom', False)
                saved = self.block_config.params.get(key, '')
                if saved in choices:
                    init_val = choice_labels[choices.index(saved)]
                elif saved and allow_custom:
                    init_val = saved
                else:
                    init_val = choice_labels[0] if choice_labels else ''
                var = tk.StringVar(value=init_val)
                values_display = choice_labels if choice_labels else choices
                mapping = dict(zip(choice_labels, choices))
                if allow_custom:
                    w = ttk.Combobox(frame, textvariable=var, values=values_display, width=28)
                else:
                    w = ttk.Combobox(frame, textvariable=var, values=values_display, width=28, state='readonly')
                w.pack(side=tk.LEFT, fill=tk.X, expand=True)
                self.param_widgets[key] = ('choice', var, w, mapping, allow_custom)
                if key == 'mode':
                    w.bind('<<ComboboxSelected>>', self._on_mode_change)

            elif ptype == 'boolean':
                var = tk.BooleanVar(value=self.block_config.params.get(key, False))
                w = ttk.Checkbutton(frame, variable=var, text="是")
                w.pack(side=tk.LEFT)
                self.param_widgets[key] = ('boolean', var)

            elif ptype == 'text':
                var = tk.StringVar(value=str(self.block_config.params.get(key, '')))
                w = ttk.Entry(frame, textvariable=var, width=32)
                w.pack(side=tk.LEFT, fill=tk.X, expand=True)
                self.param_widgets[key] = ('text', var)

            elif ptype == 'chain_list':
                chain_frame = ttk.Frame(frame)
                chain_frame.pack(fill=tk.X)
                self.chain_frame = chain_frame
                self.chain_rows = []
                chains = self.block_config.params.get(key, [])
                if not chains:
                    chains = [{}]
                for chain in chains:
                    self._add_chain_row(chain)
                ttk.Button(chain_frame, text="+ 添加查找步骤", command=lambda: self._add_chain_row({})).pack(anchor=tk.W, pady=2)
                self.param_widgets[key] = ('chain_list',)

            elif ptype == 'item_mappings':
                mapping_frame = ttk.Frame(frame)
                mapping_frame.pack(fill=tk.X)
                self.mapping_frame = mapping_frame
                self.mapping_rows = []
                mappings = self.block_config.params.get(key, [])
                if not mappings:
                    mappings = [{}]
                for m in mappings:
                    self._add_mapping_row(m)
                ttk.Button(mapping_frame, text="+ 添加映射行", command=lambda: self._add_mapping_row({})).pack(anchor=tk.W, pady=2)
                self.param_widgets[key] = ('item_mappings',)

            row += 1
        self._inner.columnconfigure(1, weight=1)

    def _find_related_datasource(self, column_key):
        for key in ('source', 'target_table'):
            if key in self.param_widgets:
                return self.param_widgets[key][1].get()
        return self.block_config.params.get('source', self.block_config.params.get('target_table', ''))

    def _on_datasource_change(self, key):
        ds_name = self.param_widgets[key][1].get()
        cols = self._get_columns_for_ds(ds_name)
        for k, info in self.param_widgets.items():
            ptype = None
            for pdef in BLOCK_PARAMS.get(self.block_config.block_type, []):
                if pdef['key'] == k:
                    ptype = pdef['type']
                    break
            if ptype == 'column':
                info[2]['values'] = cols

    def _on_mode_change(self, event=None):
        saved = self.collect_params()
        self.block_config.params = saved
        canvas = self.app.steps_canvas if self.app else None
        scroll_pos = canvas.yview()[0] if canvas else 0
        self._build_fields()
        if canvas:
            canvas.update_idletasks()
            canvas.yview_moveto(scroll_pos)

    def _add_chain_row(self, chain_data):
        row_frame = ttk.Frame(self.chain_frame)
        row_frame.pack(fill=tk.X, pady=1)

        ds_names = self._get_ds_names()

        ttk.Label(row_frame, text="表:").pack(side=tk.LEFT)
        table_var = tk.StringVar(value=chain_data.get('target_table', ''))
        table_combo = ttk.Combobox(row_frame, textvariable=table_var, values=ds_names, width=12)
        table_combo.pack(side=tk.LEFT, padx=2)

        ttk.Label(row_frame, text="查找列:").pack(side=tk.LEFT)
        find_var = tk.StringVar(value=chain_data.get('find_column', ''))
        find_combo = ttk.Combobox(row_frame, textvariable=find_var, values=[], width=12)
        find_combo.pack(side=tk.LEFT, padx=2)

        ttk.Label(row_frame, text="返回列:").pack(side=tk.LEFT)
        ret_var = tk.StringVar(value=chain_data.get('return_column', ''))
        ret_combo = ttk.Combobox(row_frame, textvariable=ret_var, values=[], width=12)
        ret_combo.pack(side=tk.LEFT, padx=2)

        def update_cols(*args):
            cols = self._get_columns_for_ds(table_var.get())
            find_combo['values'] = cols
            ret_combo['values'] = cols

        table_combo.bind('<<ComboboxSelected>>', update_cols)
        update_cols()

        ttk.Button(row_frame, text="✕", width=2,
                   command=lambda: self._remove_row(row_frame, self.chain_rows, chain_data_item)).pack(side=tk.LEFT, padx=2)

        chain_data_item = {'target_table': table_var, 'find_column': find_var, 'return_column': ret_var}
        self.chain_rows.append((row_frame, chain_data_item))

    def _add_mapping_row(self, mapping_data):
        row_frame = ttk.Frame(self.mapping_frame)
        row_frame.pack(fill=tk.X, pady=1)

        ttk.Label(row_frame, text="名称:").pack(side=tk.LEFT)
        name_var = tk.StringVar(value=mapping_data.get('name', ''))
        ttk.Entry(row_frame, textvariable=name_var, width=8).pack(side=tk.LEFT, padx=2)

        ttk.Label(row_frame, text="道具ID:").pack(side=tk.LEFT)
        id_var = tk.StringVar(value=str(mapping_data.get('item_id', '')))
        ttk.Entry(row_frame, textvariable=id_var, width=8).pack(side=tk.LEFT, padx=2)

        ttk.Label(row_frame, text="期望变量:").pack(side=tk.LEFT)
        exp_var = tk.StringVar(value=mapping_data.get('expected_var', ''))
        exp_combo = ttk.Combobox(row_frame, textvariable=exp_var, values=self._get_var_names(), width=12)
        exp_combo.pack(side=tk.LEFT, padx=2)

        ttk.Button(row_frame, text="✕", width=2,
                   command=lambda: self._remove_row(row_frame, self.mapping_rows, mapping_item)).pack(side=tk.LEFT, padx=2)

        mapping_item = {'name': name_var, 'item_id': id_var, 'expected_var': exp_var}
        self.mapping_rows.append((row_frame, mapping_item))

    def _remove_row(self, frame, row_list, item):
        frame.destroy()
        if (frame, item) in row_list:
            row_list.remove((frame, item))

    def collect_params(self):
        params = {}
        for key, widget_info in self.param_widgets.items():
            if widget_info[0] == 'combobox':
                params[key] = widget_info[1].get()
            elif widget_info[0] == 'choice':
                label_val = widget_info[1].get()
                mapping = widget_info[3]
                params[key] = mapping.get(label_val, label_val)
            elif widget_info[0] == 'boolean':
                params[key] = widget_info[1].get()
            elif widget_info[0] == 'text':
                params[key] = widget_info[1].get()
            elif widget_info[0] == 'chain_list':
                chains = []
                for _, item in self.chain_rows:
                    chains.append({
                        'target_table': item['target_table'].get(),
                        'find_column': item['find_column'].get(),
                        'return_column': item['return_column'].get(),
                    })
                params[key] = chains
            elif widget_info[0] == 'item_mappings':
                mappings = []
                for _, item in self.mapping_rows:
                    mappings.append({
                        'name': item['name'].get(),
                        'item_id': item['item_id'].get(),
                        'expected_var': item['expected_var'].get(),
                    })
                params[key] = mappings
        return params

    def update_dropdowns(self):
        for key, widget_info in self.param_widgets.items():
            if widget_info[0] != 'combobox':
                continue
            ptype = None
            for pdef in BLOCK_PARAMS.get(self.block_config.block_type, []):
                if pdef['key'] == key:
                    ptype = pdef['type']
                    break
            if ptype == 'datasource':
                widget_info[2]['values'] = self._get_ds_names()
            elif ptype == 'variable':
                widget_info[2]['values'] = self._get_var_names()
            elif ptype == 'column':
                ds_name = self._find_related_datasource(key)
                widget_info[2]['values'] = self._get_columns_for_ds(ds_name)

    def validate(self):
        errors = []
        param_defs = BLOCK_PARAMS.get(self.block_config.block_type, [])
        current_mode = self.block_config.params.get('mode', '')
        if not current_mode:
            for p in param_defs:
                if p['key'] == 'mode' and p['type'] == 'choice':
                    current_mode = p.get('choices', [''])[0]
                    break
        for pdef in param_defs:
            key = pdef['key']
            ptype = pdef['type']
            condition = pdef.get('condition', None)
            optional = pdef.get('optional', False)
            if optional:
                continue
            if condition:
                cond_key, cond_val = condition.split('==')
                check_val = current_mode if cond_key.strip() == 'mode' else self.block_config.params.get(cond_key.strip(), '')
                if str(check_val).strip() != cond_val.strip():
                    continue
            val = self.block_config.params.get(key, '')
            if ptype in ('datasource', 'variable', 'column', 'text') and not str(val).strip():
                errors.append(f"{pdef['label']} 未填写")
        return errors


class TableVerifyApp:
    @staticmethod
    def _find_base_dir():
        path = os.path.dirname(os.path.abspath(__file__))
        for _ in range(10):
            if os.path.isdir(os.path.join(path, 'data', 'apps')):
                return path
            parent = os.path.dirname(path)
            if parent == path:
                break
            path = parent
        return os.path.dirname(os.path.abspath(__file__))

    _BASE_DIR = _find_base_dir.__func__()
    SETTINGS_FILE = os.path.join(_BASE_DIR, 'settings.json')

    def __init__(self, root):
        self.root = root
        self.root.title(f"对表工具 TableVerify v{VERSION} - by.{AUTHOR}")
        self.root.minsize(900, 600)

        self.current_task = Task()
        self.current_task_file = None
        self.runner = Runner()
        self.last_result = None
        self.step_widgets = []
        self.table_cache = {}
        self._file_cache = {}
        self._widgets_group_name = None
        self._last_sash_h = 300
        self._last_sash_v = 400
        self._closing = False
        self._task_snapshot = None

        self._build_menu()
        self._build_task_toolbar()
        self._build_ui()
        self._restore_window_state()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self.root.after(500, self._track_sash_positions)
        self._new_task()

    def _mark_clean(self):
        self._collect_step_params()
        self._task_snapshot = json.dumps(self.current_task.to_dict(), ensure_ascii=False, sort_keys=True)

    def _is_dirty(self):
        self._collect_step_params()
        current = json.dumps(self.current_task.to_dict(), ensure_ascii=False, sort_keys=True)
        return current != self._task_snapshot

    def _check_unsaved(self):
        if self._is_dirty():
            result = messagebox.askyesnocancel(
                "未保存的修改",
                "当前任务有未保存的修改，是否保存？\n\n是：保存后继续\n否：不保存，直接继续\n取消：返回继续编辑",
                parent=self.root
            )
            if result is None:
                return False
            if result:
                self._save_task()
            if not self.current_task_file:
                return False
        return True

    def _build_menu(self):
        menubar = tk.Menu(self.root)
        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="新建任务", command=self._new_task, accelerator="Ctrl+N")
        file_menu.add_command(label="保存任务", command=self._save_task, accelerator="Ctrl+S")
        file_menu.add_command(label="另存为...", command=self._copy_task, accelerator="Ctrl+Shift+S")
        file_menu.add_separator()
        file_menu.add_command(label="退出", command=self.root.quit)
        menubar.add_cascade(label="文件", menu=file_menu)

        run_menu = tk.Menu(menubar, tearoff=0)
        run_menu.add_command(label="运行验证", command=self._run_verify, accelerator="F5")
        run_menu.add_command(label="批量运行...", command=self._batch_run)
        run_menu.add_command(label="导出报告...", command=self._export_report)
        menubar.add_cascade(label="运行", menu=run_menu)

        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label="关于", command=self._show_about)
        menubar.add_cascade(label="帮助", menu=help_menu)

        self.root.config(menu=menubar)
        self.root.bind('<Control-n>', lambda e: self._new_task())
        self.root.bind('<Control-s>', lambda e: self._save_task())
        self.root.bind('<Control-Shift-S>', lambda e: self._copy_task())
        self.root.bind('<F5>', lambda e: self._run_verify())

    def _show_about(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("关于")
        dlg.resizable(False, False)
        dlg.transient(self.root)
        dlg.grab_set()
        frame = ttk.Frame(dlg, padding=20)
        frame.pack()
        ttk.Label(frame, text="对表工具 TableVerify", font=('TkDefaultFont', 14, 'bold')).pack(pady=(0, 5))
        ttk.Label(frame, text=f"版本: v{VERSION}", font=('TkDefaultFont', 10)).pack()
        ttk.Label(frame, text=f"作者: {AUTHOR}", font=('TkDefaultFont', 10)).pack(pady=(5, 15))
        ttk.Button(frame, text="确定", command=dlg.destroy, width=8).pack()
        dlg.geometry("+%d+%d" % (self.root.winfo_x() + 200, self.root.winfo_y() + 200))

    def _build_task_toolbar(self):
        toolbar = ttk.Frame(self.root, padding=(5, 4))
        toolbar.pack(side=tk.TOP, fill=tk.X)

        ttk.Label(toolbar, text="当前任务:", font=('TkDefaultFont', 10)).pack(side=tk.LEFT, padx=(0, 4))
        self.task_var = tk.StringVar()
        self.task_combo = ttk.Combobox(toolbar, textvariable=self.task_var, width=30, state='readonly')
        self.task_combo.pack(side=tk.LEFT, padx=2)
        self.task_combo.bind('<<ComboboxSelected>>', self._on_task_selected)

        ttk.Button(toolbar, text="新建", command=self._new_task, width=6).pack(side=tk.LEFT, padx=3)
        ttk.Button(toolbar, text="保存", command=self._save_task, width=6).pack(side=tk.LEFT, padx=3)
        ttk.Button(toolbar, text="另存为", command=self._copy_task, width=6).pack(side=tk.LEFT, padx=3)
        ttk.Button(toolbar, text="删除", command=self._delete_task, width=6).pack(side=tk.LEFT, padx=3)
        ttk.Button(toolbar, text="重命名", command=self._rename_task, width=6).pack(side=tk.LEFT, padx=3)
        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=10)
        ttk.Button(toolbar, text="运行 (F5)", command=self._run_verify, width=10).pack(side=tk.LEFT, padx=3)
        ttk.Button(toolbar, text="批量运行", command=self._batch_run, width=8).pack(side=tk.LEFT, padx=3)
        ttk.Button(toolbar, text="导出报告", command=self._export_report, width=8).pack(side=tk.LEFT, padx=3)
        self._refresh_task_list()

    def _refresh_task_list(self):
        tasks_dir = os.path.join(self._BASE_DIR, 'Tasks')
        os.makedirs(tasks_dir, exist_ok=True)
        task_files = [f[:-5] for f in os.listdir(tasks_dir) if f.endswith('.task')]
        self.task_combo['values'] = task_files
        if self.current_task.name in task_files:
            self.task_combo.set(self.current_task.name)
        elif task_files:
            self.task_combo.set('')

    def _on_task_selected(self, event=None):
        name = self.task_var.get()
        if not name or (self.current_task.name == name and self.current_task_file):
            return
        if not self._check_unsaved():
            if self.current_task_file:
                self.task_combo.set(os.path.basename(self.current_task_file).replace('.task', ''))
            return
        tasks_dir = os.path.join(self._BASE_DIR, 'Tasks')
        file_path = os.path.join(tasks_dir, f"{name}.task")
        if os.path.exists(file_path):
            try:
                self.status_bar.config(text=f"正在加载: {name}...")
                self.root.update()
                self.current_task = Task.load(file_path)
                self.current_task_file = file_path
                self._refresh_all()
                self.task_combo.set(name)
                self._mark_clean()
                self.status_bar.config(text=f"已加载任务: {name}")
            except Exception as e:
                messagebox.showerror("错误", f"加载任务失败: {e}")

    def _build_ui(self):
        main_area = ttk.Frame(self.root)
        main_area.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=2, pady=(0, 2))

        self.paned_h = tk.PanedWindow(main_area, orient=tk.HORIZONTAL, sashwidth=4, bg='#d0d0d0')
        self.paned_h.pack(fill=tk.BOTH, expand=True)

        left_frame = ttk.LabelFrame(self.paned_h, text="数据源管理")
        self.paned_h.add(left_frame, minsize=100)
        right_frame = ttk.Frame(self.paned_h)
        self.paned_h.add(right_frame, minsize=400)

        self._build_datasource_panel(left_frame)

        self.paned_v = tk.PanedWindow(right_frame, orient=tk.VERTICAL, sashwidth=4, bg='#d0d0d0')
        self.paned_v.pack(fill=tk.BOTH, expand=True)
        steps_frame = ttk.LabelFrame(self.paned_v, text="验证步骤")
        self.paned_v.add(steps_frame, minsize=100)
        self._build_steps_panel(steps_frame)
        results_frame = ttk.LabelFrame(self.paned_v, text="运行结果")
        self.paned_v.add(results_frame, minsize=100)
        self._build_results_panel(results_frame)

        self.status_bar = ttk.Label(self.root, text="就绪", relief=tk.SUNKEN, anchor=tk.W, padding=(5, 2))
        self.status_bar.pack(side=tk.BOTTOM, fill=tk.X)

    def _build_datasource_panel(self, parent):
        toolbar = ttk.Frame(parent)
        toolbar.pack(fill=tk.X, padx=5, pady=3)
        ttk.Button(toolbar, text="+ 添加", command=self._add_datasource).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="编辑", command=self._edit_datasource).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="删除", command=self._remove_datasource).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="预览", command=self._preview_datasource).pack(side=tk.LEFT, padx=2)

        self.ds_listbox = tk.Listbox(parent, height=15, font=("Consolas", 10))
        self.ds_listbox.pack(fill=tk.BOTH, expand=True, padx=5, pady=3)

        self.ds_info = ttk.Label(parent, text="", foreground='gray', wraplength=260, justify=tk.LEFT)
        self.ds_info.pack(fill=tk.X, padx=5, pady=3)
        self.ds_listbox.bind('<<ListboxSelect>>', self._on_ds_select)

    def _build_steps_panel(self, parent):
        group_toolbar = ttk.Frame(parent)
        group_toolbar.pack(fill=tk.X, padx=5, pady=3)

        ttk.Label(group_toolbar, text="验证组:").pack(side=tk.LEFT, padx=2)
        self.group_var = tk.StringVar()
        self.group_combo = ttk.Combobox(group_toolbar, textvariable=self.group_var, width=20, state='readonly')
        self.group_combo.pack(side=tk.LEFT, padx=2)
        self.group_combo.bind('<<ComboboxSelected>>', self._on_group_change)

        ttk.Button(group_toolbar, text="+ 新建组", command=self._add_group).pack(side=tk.LEFT, padx=5)
        ttk.Button(group_toolbar, text="删除组", command=self._remove_group).pack(side=tk.LEFT, padx=2)
        ttk.Separator(group_toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=8)

        self.add_step_menu = tk.Menu(group_toolbar, tearoff=0)
        for btype, bname in BLOCK_TYPES.items():
            self.add_step_menu.add_command(
                label=f"{bname} - {BLOCK_DESCRIPTIONS.get(btype, '')}",
                command=lambda t=btype: self._add_step_type(t)
            )
        btn_add = ttk.Menubutton(group_toolbar, text="+ 添加步骤", menu=self.add_step_menu)
        btn_add.pack(side=tk.LEFT, padx=2)

        canvas_frame = ttk.Frame(parent)
        canvas_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=3)

        self.steps_canvas = tk.Canvas(canvas_frame, highlightthickness=0)
        scrollbar = ttk.Scrollbar(canvas_frame, orient=tk.VERTICAL, command=self.steps_canvas.yview)
        self.steps_inner = ttk.Frame(self.steps_canvas)

        self._canvas_window = self.steps_canvas.create_window((0, 0), window=self.steps_inner, anchor=tk.NW)
        self.steps_canvas.configure(yscrollcommand=scrollbar.set)

        self.steps_inner.bind('<Configure>', lambda e: self.steps_canvas.configure(scrollregion=self.steps_canvas.bbox('all')))
        self.steps_canvas.bind('<Configure>', lambda e: self.steps_canvas.itemconfig(self._canvas_window, width=e.width))

        self.steps_canvas.bind('<MouseWheel>', self._on_steps_mousewheel)
        self.steps_canvas.bind('<Button-4>', self._on_steps_mousewheel)
        self.steps_canvas.bind('<Button-5>', self._on_steps_mousewheel)
        self.steps_canvas.bind('<Enter>', self._bind_steps_mousewheel)
        self.steps_canvas.bind('<Leave>', self._unbind_steps_mousewheel)

        self.steps_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def _bind_steps_mousewheel(self, event=None):
        self._bind_mousewheel_to_children(self.steps_inner)

    def _unbind_steps_mousewheel(self, event=None):
        self._unbind_mousewheel_from_children(self.steps_inner)

    def _bind_mousewheel_to_children(self, widget):
        if not isinstance(widget, (ttk.Combobox, tk.Listbox, ttk.Entry, tk.Entry, tk.Spinbox, ttk.Spinbox)):
            widget.bind('<MouseWheel>', self._on_steps_mousewheel)
            widget.bind('<Button-4>', self._on_steps_mousewheel)
            widget.bind('<Button-5>', self._on_steps_mousewheel)
        for child in widget.winfo_children():
            self._bind_mousewheel_to_children(child)

    def _unbind_mousewheel_from_children(self, widget):
        try:
            widget.unbind('<MouseWheel>')
            widget.unbind('<Button-4>')
            widget.unbind('<Button-5>')
        except Exception:
            pass
        for child in widget.winfo_children():
            self._unbind_mousewheel_from_children(child)

    def _on_steps_mousewheel(self, event):
        if event.num == 4:
            self.steps_canvas.yview_scroll(-3, 'units')
        elif event.num == 5:
            self.steps_canvas.yview_scroll(3, 'units')
        else:
            self.steps_canvas.yview_scroll(int(-1 * (event.delta / 120)), 'units')

    def _build_results_panel(self, parent):
        cols = ('label', 'item_name', 'expected', 'actual', 'result', 'message')
        self.result_tree = ttk.Treeview(parent, columns=cols, show='headings', height=10)
        for c, t in zip(cols, ('标签', '道具', '期望值', '实际值', '结果', '说明')):
            self.result_tree.heading(c, text=t)
            self.result_tree.column(c, width=80, minwidth=50, stretch=True)

        vsb = ttk.Scrollbar(parent, orient=tk.VERTICAL, command=self.result_tree.yview)
        hsb = ttk.Scrollbar(parent, orient=tk.HORIZONTAL, command=self.result_tree.xview)
        self.result_tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        hsb.pack(side=tk.BOTTOM, fill=tk.X)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)
        self.result_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self.result_tree.tag_configure('pass', foreground='#2e7d32')
        self.result_tree.tag_configure('fail', foreground='#c62828')
        self.result_tree.tag_configure('skip', foreground='#9e9e9e')
        self.result_tree.tag_configure('group_header', background='#e3f2fd', font=('TkDefaultFont', 10, 'bold'))
        self.result_tree.tag_configure('task_header', background='#c8e6c9', font=('TkDefaultFont', 10, 'bold'))

    def _refresh_table_cache(self):
        self.table_cache.clear()
        for ds in self.current_task.data_sources:
            try:
                self.table_cache[ds.name] = self._load_table_cached(ds)
            except Exception:
                pass

    def _load_table_cached(self, ds):
        cache_key = f"{ds.file_path}|{ds.sheet_name}|{ds.header_row}"
        try:
            mtime = str(os.path.getmtime(ds.file_path))
        except Exception:
            mtime = ''
        cached = self._file_cache.get(cache_key)
        if cached and cached[0] == mtime:
            return cached[1]
        table = TableLoader.load(name=ds.name, file_path=ds.file_path, sheet_name=ds.sheet_name, header_row=ds.header_row)
        self._file_cache[cache_key] = (mtime, table)
        return table

    def _on_close(self):
        if self._is_dirty():
            result = messagebox.askyesnocancel(
                "未保存的修改",
                "当前任务有未保存的修改。\n\n是：保存后关闭\n否：不保存，直接关闭\n取消：返回继续编辑",
                parent=self.root
            )
            if result is None:
                return
            if result:
                self._collect_step_params()
                self._save_task()
                if not self.current_task_file:
                    return
        self._closing = True
        self._save_window_state()
        self.root.destroy()

    def _track_sash_positions(self):
        if self._closing:
            return
        try:
            pos_h = self.paned_h.sashpos(0)
            pos_v = self.paned_v.sashpos(0)
            if pos_h > 50:
                self._last_sash_h = pos_h
            if pos_v > 50:
                self._last_sash_v = pos_v
        except Exception:
            pass
        self.root.after(500, self._track_sash_positions)

    def _save_window_state(self):
        try:
            geo = self.root.geometry()
            last_task = os.path.basename(self.current_task_file) if self.current_task_file else ''
            settings = {
                'geometry': geo,
                'paned_h_sash': self._last_sash_h,
                'paned_v_sash': self._last_sash_v,
                'last_task': last_task,
            }
            with open(self.SETTINGS_FILE, 'w', encoding='utf-8') as f:
                json.dump(settings, f)
        except Exception:
            pass

    def _restore_window_state(self):
        self.root.geometry("1280x800")
        try:
            if not os.path.exists(self.SETTINGS_FILE):
                return
            with open(self.SETTINGS_FILE, 'r', encoding='utf-8') as f:
                settings = json.load(f)
            geo = settings.get('geometry', '1280x800')
            self.root.geometry(geo)
            self.root.update_idletasks()
            paned_h_sash = settings.get('paned_h_sash', 0)
            paned_v_sash = settings.get('paned_v_sash', 0)
            if paned_h_sash > 50:
                self._last_sash_h = paned_h_sash
            if paned_v_sash > 50:
                self._last_sash_v = paned_v_sash
            self.root.after(200, self._apply_sash_positions)
            last_task = settings.get('last_task', '')
            if last_task:
                tasks_dir = os.path.join(self._BASE_DIR, 'Tasks')
                task_path = os.path.join(tasks_dir, last_task)
                if os.path.exists(task_path):
                    self.root.after(400, lambda: self._load_task_file(task_path))
        except Exception:
            pass

    def _apply_sash_positions(self):
        try:
            self.paned_h.sashpos(0, self._last_sash_h)
            self.paned_v.sashpos(0, self._last_sash_v)
        except Exception:
            pass

    def _load_task_file(self, file_path):
        try:
            name = os.path.basename(file_path).replace('.task', '')
            self.status_bar.config(text=f"正在加载: {name}...")
            self.root.update()
            self.current_task = Task.load(file_path)
            self.current_task_file = file_path
            self._refresh_all()
            self.task_combo.set(self.current_task.name)
            self._mark_clean()
            self.status_bar.config(text=f"已加载任务: {self.current_task.name}")
        except Exception:
            pass

    def _new_task(self):
        if self.step_widgets and not self._check_unsaved():
            return
        self.current_task = Task(name="新任务")
        self.current_task_file = None
        self._refresh_all()
        self._mark_clean()

    def _refresh_all(self):
        self._refresh_table_cache()
        self._refresh_task_list()
        self._refresh_ds_list()
        self._refresh_group_combo()
        self._refresh_steps()
        self._clear_results()

    def _refresh_ds_list(self):
        self.ds_listbox.delete(0, tk.END)
        for ds in self.current_task.data_sources:
            ext = os.path.splitext(ds.file_path)[1].upper().lstrip('.')
            sheet_info = f" / {ds.sheet_name}" if ds.sheet_name else ""
            self.ds_listbox.insert(tk.END, f"[{ext}{sheet_info}] {ds.name}")

    def _refresh_group_combo(self):
        names = [g.name for g in self.current_task.groups]
        self.group_combo['values'] = names
        if names and self.group_var.get() not in names:
            self.group_combo.set(names[0])
        elif not names:
            self.group_var.set('')

    def _on_ds_select(self, event=None):
        sel = self.ds_listbox.curselection()
        if sel:
            ds = self.current_task.data_sources[sel[0]]
            table = self.table_cache.get(ds.name)
            col_count = len(table.headers) if table else '?'
            row_count = len(table.raw_data) - table.header_row if table else '?'
            self.ds_info.config(text=f"路径: {ds.file_path}\n表头行: {ds.header_row}  列数: {col_count}  数据行: {row_count}")

    def _add_datasource(self):
        dlg = DataSourceDialog(self.root)
        self.root.wait_window(dlg)
        if dlg.result:
            if any(ds.name == dlg.result.name for ds in self.current_task.data_sources):
                messagebox.showwarning("提示", f"数据源名称 '{dlg.result.name}' 已存在")
                return
            self.current_task.data_sources.append(dlg.result)
            self._refresh_table_cache()
            self._refresh_ds_list()
            self._update_step_dropdowns()

    def _edit_datasource(self):
        sel = self.ds_listbox.curselection()
        if not sel:
            return
        idx = sel[0]
        dlg = DataSourceDialog(self.root, self.current_task.data_sources[idx])
        self.root.wait_window(dlg)
        if dlg.result:
            self.current_task.data_sources[idx] = dlg.result
            self._refresh_table_cache()
            self._refresh_ds_list()
            self._update_step_dropdowns()

    def _remove_datasource(self):
        sel = self.ds_listbox.curselection()
        if not sel:
            return
        name = self.current_task.data_sources[sel[0]].name
        if messagebox.askyesno("确认", f"确定删除数据源 '{name}'?"):
            self.current_task.data_sources.pop(sel[0])
            self._refresh_table_cache()
            self._refresh_ds_list()
            self._update_step_dropdowns()

    def _preview_datasource(self):
        sel = self.ds_listbox.curselection()
        if not sel:
            messagebox.showwarning("提示", "请先选择一个数据源")
            return
        ds = self.current_task.data_sources[sel[0]]
        DataPreviewDialog(self.root, ds, self.table_cache)

    def _update_step_dropdowns(self):
        for sw in self.step_widgets:
            sw.update_dropdowns()

    def _add_group(self):
        name = simpledialog.askstring("新建验证组", "验证组名称:", parent=self.root)
        if name:
            self.current_task.groups.append(VerifyGroup(name=name))
            self._refresh_group_combo()
            self.group_combo.set(name)
            self._refresh_steps()

    def _remove_group(self):
        name = self.group_var.get()
        if not name:
            return
        if messagebox.askyesno("确认", f"确定删除验证组 '{name}'?"):
            self.current_task.groups = [g for g in self.current_task.groups if g.name != name]
            self._refresh_group_combo()
            if self.current_task.groups:
                self.group_combo.set(self.current_task.groups[0].name)
            self._refresh_steps()

    def _on_group_change(self, event=None):
        self._refresh_steps()

    def _get_current_group(self):
        name = self.group_var.get()
        for g in self.current_task.groups:
            if g.name == name:
                return g
        return None

    def _add_step_type(self, block_type):
        group = self._get_current_group()
        if not group:
            messagebox.showwarning("提示", "请先创建验证组")
            return
        group.steps.append(BlockConfig(block_type=block_type))
        self._refresh_steps()

    def _copy_step(self, idx):
        self._collect_step_params()
        group = self._get_current_group()
        if not group or idx >= len(group.steps):
            return
        new_step = copy.deepcopy(group.steps[idx])
        group.steps.insert(idx + 1, new_step)
        self._refresh_steps()

    def _toggle_step(self, idx):
        group = self._get_current_group()
        if not group or idx >= len(group.steps):
            return
        step = group.steps[idx]
        step.params['_disabled'] = not step.params.get('_disabled', False)
        self._refresh_steps()

    def _refresh_steps(self):
        self._collect_step_params()
        for w in self.steps_inner.winfo_children():
            w.destroy()
        self.step_widgets = []
        self._widgets_group_name = None

        group = self._get_current_group()
        if not group:
            return

        self._widgets_group_name = group.name

        for idx, step in enumerate(group.steps):
            disabled = step.params.get('_disabled', False)
            step_frame = ttk.Frame(self.steps_inner)
            step_frame.pack(fill=tk.X, pady=3, padx=2)

            header = ttk.Frame(step_frame)
            header.pack(fill=tk.X)

            btype_name = BLOCK_TYPES.get(step.block_type, step.block_type)
            state_tag = " [已禁用]" if disabled else ""
            lbl = ttk.Label(header, text=f"步骤{idx + 1}: {btype_name}{state_tag}", font=('TkDefaultFont', 10, 'bold'))
            lbl.pack(side=tk.LEFT, padx=5)
            if disabled:
                lbl.configure(foreground='gray')

            btn_frame = ttk.Frame(header)
            btn_frame.pack(side=tk.RIGHT, padx=5)
            ttk.Button(btn_frame, text="复制", command=lambda i=idx: self._copy_step(i)).pack(side=tk.LEFT, padx=1)
            ttk.Button(btn_frame, text="禁用" if not disabled else "启用", command=lambda i=idx: self._toggle_step(i)).pack(side=tk.LEFT, padx=1)
            ttk.Button(btn_frame, text="▲", width=2, command=lambda i=idx: self._move_step(i, -1)).pack(side=tk.LEFT, padx=1)
            ttk.Button(btn_frame, text="▼", width=2, command=lambda i=idx: self._move_step(i, 1)).pack(side=tk.LEFT, padx=1)
            ttk.Button(btn_frame, text="✕", width=2, command=lambda i=idx: self._remove_step(i)).pack(side=tk.LEFT, padx=1)

            sw = StepWidget(step_frame, step, self, step_index=idx)
            if disabled:
                sw.configure(foreground='gray')
            sw.pack(fill=tk.X, padx=5, pady=2)
            self.step_widgets.append(sw)

        self._bind_mousewheel_to_children(self.steps_inner)

    def _move_step(self, idx, direction):
        group = self._get_current_group()
        if not group:
            return
        new_idx = idx + direction
        if 0 <= new_idx < len(group.steps):
            group.steps[idx], group.steps[new_idx] = group.steps[new_idx], group.steps[idx]
            self._refresh_steps()

    def _remove_step(self, idx):
        group = self._get_current_group()
        if not group:
            return
        group.steps.pop(idx)
        self._refresh_steps()

    def _collect_step_params(self):
        if not self._widgets_group_name:
            return
        group = None
        for g in self.current_task.groups:
            if g.name == self._widgets_group_name:
                group = g
                break
        if not group:
            return
        for idx, sw in enumerate(self.step_widgets):
            if idx < len(group.steps):
                group.steps[idx].params = sw.collect_params()

    def _validate_before_run(self):
        self._collect_step_params()
        errors = []
        if not self.current_task.groups:
            return ["没有验证组"]
        for group in self.current_task.groups:
            if not group.steps:
                errors.append(f"验证组 '{group.name}' 中没有步骤")
                continue
            for idx, step in enumerate(group.steps):
                if step.params.get('_disabled', False):
                    continue
                sw_idx = -1
                if group.name == self._widgets_group_name:
                    active_idx = 0
                    for si, s in enumerate(group.steps):
                        if not s.params.get('_disabled', False):
                            if active_idx < len(self.step_widgets):
                                sw = self.step_widgets[active_idx]
                                sw_idx = si
                            active_idx += 1
                    sw = None
                else:
                    sw = None
                step_errors = []
                for pdef in BLOCK_PARAMS.get(step.block_type, []):
                    key = pdef['key']
                    ptype = pdef['type']
                    condition = pdef.get('condition', None)
                    optional = pdef.get('optional', False)
                    if optional:
                        continue
                    current_mode = step.params.get('mode', '')
                    if not current_mode:
                        for p in BLOCK_PARAMS.get(step.block_type, []):
                            if p['key'] == 'mode' and p['type'] == 'choice':
                                current_mode = p.get('choices', [''])[0]
                                break
                    if condition:
                        cond_key, cond_val = condition.split('==')
                        check_val = current_mode if cond_key.strip() == 'mode' else step.params.get(cond_key.strip(), '')
                        if str(check_val).strip() != cond_val.strip():
                            continue
                    val = step.params.get(key, '')
                    if ptype in ('datasource', 'variable', 'column', 'text') and not str(val).strip():
                        step_errors.append(f"{pdef['label']} 未填写")
                for err in step_errors:
                    errors.append(f"[{group.name}] 步骤{idx + 1}: {err}")
        return errors

    def _run_verify(self):
        errors = self._validate_before_run()
        if errors:
            messagebox.showwarning("参数不完整", "以下参数未填写:\n\n" + "\n".join(errors[:10]))
            return

        self.status_bar.config(text="正在运行验证...")
        self.root.update()

        def run():
            try:
                filtered_task = copy.deepcopy(self.current_task)
                for g in filtered_task.groups:
                    g.steps = [s for s in g.steps if not s.params.get('_disabled', False)]
                    for s in g.steps:
                        s.params.pop('_disabled', None)
                self.runner.tables = {}
                for ds in filtered_task.data_sources:
                    try:
                        self.runner.tables[ds.name] = self._load_table_cached(ds)
                    except Exception:
                        pass
                result = self.runner.run(filtered_task, skip_load=True)
                self.last_result = [result]
                self.root.after(0, lambda: self._show_results([result]))
            except Exception as e:
                self.root.after(0, lambda: self.status_bar.config(text=f"运行错误: {e}"))

        thread = threading.Thread(target=run, daemon=True)
        thread.start()

    def _batch_run(self):
        tasks_dir = os.path.join(self._BASE_DIR, 'Tasks')
        os.makedirs(tasks_dir, exist_ok=True)
        task_files = [f[:-5] for f in os.listdir(tasks_dir) if f.endswith('.task')]
        if not task_files:
            messagebox.showwarning("提示", "没有已保存的任务可供运行", parent=self.root)
            return

        dlg = BatchRunDialog(self.root, task_files)
        self.root.wait_window(dlg)
        if not dlg.result:
            return

        selected_names = dlg.result
        self.status_bar.config(text=f"正在批量运行 (0/{len(selected_names)})...")
        self.root.update()

        def run():
            all_results = []
            for i, name in enumerate(selected_names):
                try:
                    self.root.after(0, lambda n=name, idx=i: self.status_bar.config(
                        text=f"正在批量运行 ({idx + 1}/{len(selected_names)}): {n}..."))
                    task_path = os.path.join(tasks_dir, f"{name}.task")
                    task = Task.load(task_path)
                    for g in task.groups:
                        g.steps = [s for s in g.steps if not s.params.get('_disabled', False)]
                        for s in g.steps:
                            s.params.pop('_disabled', None)
                    runner = Runner()
                    runner.tables = {}
                    for ds in task.data_sources:
                        try:
                            runner.tables[ds.name] = self._load_table_cached(ds)
                        except Exception:
                            pass
                    result = runner.run(task, skip_load=True)
                    all_results.append(result)
                except Exception as e:
                    r = RunResult(name)
                    gr = GroupResult("加载错误")
                    gr.error = str(e)
                    r.group_results.append(gr)
                    all_results.append(r)

            self.last_result = all_results
            self.root.after(0, lambda: self._show_results(all_results))

        thread = threading.Thread(target=run, daemon=True)
        thread.start()

    def _show_results(self, results):
        self._clear_results()
        if not isinstance(results, list):
            results = [results]
        total_p, total_f, total_s = 0, 0, 0
        for result in results:
            tp, tf, ts = result.total_passed, result.total_failed, result.total_skipped
            total_p += tp
            total_f += tf
            total_s += ts
            summary = f"通过:{tp} 失败:{tf} 跳过:{ts}"
            task_id = self.result_tree.insert('', tk.END,
                values=(f"[任务: {result.task_name}]", '', '', '', summary, ''),
                tags=('task_header',))
            for gr in result.group_results:
                group_id = self.result_tree.insert(task_id, tk.END,
                    values=(f"[{gr.group_name}]", '', '', '', '', ''),
                    tags=('group_header',))
                if gr.error:
                    self.result_tree.insert(group_id, tk.END,
                        values=('', '', '', '错误', gr.error, ''), tags=('fail',))
                    continue
                for d in gr.details:
                    tag = 'skip' if d.skipped else ('pass' if d.passed else 'fail')
                    status = '跳过' if d.skipped else ('通过' if d.passed else '失败')
                    self.result_tree.insert(group_id, tk.END, values=(
                        d.label, d.item_name or '',
                        str(d.expected) if d.expected is not None else '',
                        str(d.actual) if d.actual is not None else '',
                        status, d.message,
                    ), tags=(tag,))

        status = f"完成! 共{len(results)}个任务 通过:{total_p} 失败:{total_f} 跳过:{total_s}"
        if total_f > 0:
            status += f"  存在{total_f}个差异"
        else:
            status += "  全部通过"
        self.status_bar.config(text=status)

    def _clear_results(self):
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)

    def _save_task(self):
        self._collect_step_params()
        if self.current_task_file and os.path.exists(self.current_task_file):
            self.current_task.save(self.current_task_file)
            self._refresh_task_list()
            self._mark_clean()
            self.status_bar.config(text=f"任务已保存: {self.current_task_file}")
        else:
            tasks_dir = os.path.join(self._BASE_DIR, 'Tasks')
            os.makedirs(tasks_dir, exist_ok=True)
            file_path = filedialog.asksaveasfilename(
                initialdir=tasks_dir, defaultextension=".task",
                filetypes=[("任务文件", "*.task")], initialfile=f"{self.current_task.name}.task")
            if file_path:
                self.current_task.save(file_path)
                self.current_task_file = file_path
                self._refresh_task_list()
                self.task_combo.set(self.current_task.name)
                self._mark_clean()
                self.status_bar.config(text=f"任务已保存: {file_path}")

    def _copy_task(self):
        self._collect_step_params()
        new_name = simpledialog.askstring("另存为新任务", f"当前任务: {self.current_task.name}\n请输入新任务名称:", parent=self.root)
        if not new_name:
            return
        task_dict = self.current_task.to_dict()
        task_dict['task_name'] = new_name
        tasks_dir = os.path.join(self._BASE_DIR, 'Tasks')
        os.makedirs(tasks_dir, exist_ok=True)
        file_path = os.path.join(tasks_dir, f"{new_name}.task")
        if os.path.exists(file_path) and not messagebox.askyesno("确认", f"任务 '{new_name}' 已存在，是否覆盖？"):
            return
        with open(file_path, 'w', encoding='utf-8') as f:
            json.dump(task_dict, f, ensure_ascii=False, indent=2)
        self.current_task = Task.from_dict(task_dict)
        self.current_task_file = file_path
        self._refresh_all()
        self._mark_clean()
        self.status_bar.config(text=f"任务已另存为: {file_path}")

    def _delete_task(self):
        if not self.current_task_file or not os.path.exists(self.current_task_file):
            messagebox.showwarning("提示", "当前任务未保存，无法删除", parent=self.root)
            return
        name = self.current_task.name
        if not messagebox.askyesno("确认删除", f"确定要删除任务 '{name}' 吗？\n\n此操作不可撤销！", parent=self.root):
            return
        try:
            os.remove(self.current_task_file)
            self.current_task = Task(name="新任务")
            self.current_task_file = None
            self._refresh_all()
            self._mark_clean()
            self.status_bar.config(text=f"已删除任务: {name}")
        except Exception as e:
            messagebox.showerror("错误", f"删除任务失败: {e}", parent=self.root)

    def _rename_task(self):
        if not self.current_task_file or not os.path.exists(self.current_task_file):
            messagebox.showwarning("提示", "当前任务未保存，无法重命名", parent=self.root)
            return
        old_name = self.current_task.name
        new_name = simpledialog.askstring("重命名任务", f"当前名称: {old_name}\n请输入新名称:", parent=self.root)
        if not new_name or new_name == old_name:
            return
        tasks_dir = os.path.join(self._BASE_DIR, 'Tasks')
        new_path = os.path.join(tasks_dir, f"{new_name}.task")
        if os.path.exists(new_path):
            messagebox.showwarning("提示", f"任务 '{new_name}' 已存在", parent=self.root)
            return
        try:
            self.current_task.name = new_name
            self.current_task.save(self.current_task_file)
            os.rename(self.current_task_file, new_path)
            self.current_task_file = new_path
            self._refresh_task_list()
            self.task_combo.set(new_name)
            self._mark_clean()
            self.status_bar.config(text=f"已重命名: {old_name} → {new_name}")
        except Exception as e:
            self.current_task.name = old_name
            messagebox.showerror("错误", f"重命名失败: {e}", parent=self.root)

    def _export_report(self):
        if not self.last_result:
            messagebox.showwarning("提示", "请先运行验证")
            return
        reports_dir = os.path.join(self._BASE_DIR, 'reports')
        os.makedirs(reports_dir, exist_ok=True)
        file_path = filedialog.asksaveasfilename(
            initialdir=reports_dir, defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")], initialfile=f"{self.current_task.name}_报告.xlsx")
        if file_path:
            try:
                self._do_export_report(file_path)
                self.status_bar.config(text=f"报告已导出: {file_path}")
                messagebox.showinfo("成功", f"报告已导出到:\n{file_path}")
            except Exception as e:
                messagebox.showerror("错误", f"导出报告失败: {e}")

    def _do_export_report(self, file_path):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        results = self.last_result if isinstance(self.last_result, list) else [self.last_result]

        wb = Workbook()
        first = True
        for result in results:
            for gr in result.group_results:
                ws = wb.active if first else wb.create_sheet()
                sheet_name = f"{result.task_name}_{gr.group_name}" if len(results) > 1 else gr.group_name
                ws.title = sheet_name[:31]
                first = False

                headers = ['标签', '道具', '期望值', '实际值', '结果', '说明']
                ws.append(headers)
                for col_idx, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                    cell.alignment = Alignment(horizontal='center')

                pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
                fail_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
                skip_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

                if gr.error:
                    ws.append(['错误', '', '', '', '失败', gr.error])
                    continue

                for d in gr.details:
                    status = '跳过' if d.skipped else ('通过' if d.passed else '失败')
                    ws.append([d.label, d.item_name or '', str(d.expected) if d.expected is not None else '', str(d.actual) if d.actual is not None else '', status, d.message])
                    fill = skip_fill if d.skipped else (pass_fill if d.passed else fail_fill)
                    for col in range(1, 7):
                        ws.cell(row=ws.max_row, column=col).fill = fill

                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        wb.save(file_path)


def main():
    root = tk.Tk()
    app = TableVerifyApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
