import csv
import os
from openpyxl import load_workbook


def col_letter_to_index(letter):
    result = 0
    for char in letter.upper():
        result = result * 26 + (ord(char) - ord('A') + 1)
    return result - 1


def index_to_col_letter(index):
    result = ""
    index += 1
    while index > 0:
        index -= 1
        result = chr(ord('A') + index % 26) + result
        index //= 26
    return result


def parse_range(range_str):
    range_str = range_str.strip()
    if ':' in range_str:
        parts = range_str.split(':')
        start_str = parts[0].strip()
        end_str = parts[1].strip()

        sc, sr = _split_cell(start_str)
        ec, er = _split_cell(end_str)

        start_col = col_letter_to_index(sc) if sc else 0
        start_row = int(sr) if sr else 1
        end_col = col_letter_to_index(ec) if ec else start_col
        end_row = int(er) if er else start_row
        return (start_col, start_row, end_col, end_row)
    else:
        c, r = _split_cell(range_str)
        col = col_letter_to_index(c) if c else 0
        row = int(r) if r else 1
        return (col, row, col, row)


def _split_cell(cell_str):
    col_part = ''
    row_part = ''
    for ch in cell_str:
        if ch.isalpha():
            col_part += ch
        elif ch.isdigit():
            row_part += ch
    return col_part, row_part


class TableData:
    def __init__(self, name, file_path):
        self.name = name
        self.file_path = file_path
        self.raw_data = []
        self.headers = []
        self.header_row = 1
        self.index = {}

    def build_index(self):
        self.index = {}
        data_start = self.header_row
        for col_idx, header in enumerate(self.headers):
            if not header or str(header).strip() == '':
                continue
            col_index = {}
            for row_idx in range(data_start, len(self.raw_data)):
                value = self._get_cell(row_idx, col_idx)
                str_val = str(value) if value != '' else ''
                if str_val not in col_index:
                    col_index[str_val] = []
                col_index[str_val].append(row_idx)
            self.index[str(header).strip()] = col_index

    def _get_cell(self, row_idx, col_idx):
        if row_idx < len(self.raw_data):
            row = self.raw_data[row_idx]
            if col_idx < len(row):
                return row[col_idx]
        return ''

    def get_range(self, range_str):
        start_col, start_row, end_col, end_row = parse_range(range_str)
        start_row_0 = start_row - 1
        end_row_0 = end_row - 1
        result = []
        for row_idx in range(start_row_0, end_row_0 + 1):
            for col_idx in range(start_col, end_col + 1):
                result.append(self._get_cell(row_idx, col_idx))
        return result

    def get_range_2d(self, range_str):
        start_col, start_row, end_col, end_row = parse_range(range_str)
        start_row_0 = start_row - 1
        end_row_0 = end_row - 1
        result = []
        for row_idx in range(start_row_0, end_row_0 + 1):
            row_data = []
            for col_idx in range(start_col, end_col + 1):
                row_data.append(self._get_cell(row_idx, col_idx))
            result.append(row_data)
        return result

    def find_rows(self, column, value):
        col_key = str(column).strip()
        if col_key not in self.index:
            return []
        str_value = str(value)
        if str_value not in self.index[col_key]:
            return []
        return [self._row_to_dict(idx) for idx in self.index[col_key][str_value]]

    def find_row(self, column, value):
        rows = self.find_rows(column, value)
        return rows[0] if rows else None

    def _row_to_dict(self, row_idx):
        result = {}
        row = self.raw_data[row_idx]
        for col_idx, header in enumerate(self.headers):
            if header and str(header).strip():
                result[str(header).strip()] = row[col_idx] if col_idx < len(row) else ''
        return result

    def get_column_values(self, column_name):
        col_key = str(column_name).strip()
        if col_key not in self.headers:
            return []
        col_idx = self.headers.index(col_key)
        data_start = self.header_row
        result = []
        for row_idx in range(data_start, len(self.raw_data)):
            result.append(self._get_cell(row_idx, col_idx))
        return result


class TableLoader:
    @staticmethod
    def load(name, file_path, sheet_name=None, header_row=2):
        ext = os.path.splitext(file_path)[1].lower()
        if ext == '.csv':
            return TableLoader._load_csv(name, file_path, header_row)
        elif ext in ('.xlsx', '.xlsm'):
            return TableLoader._load_excel(name, file_path, sheet_name, header_row)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

    @staticmethod
    def _load_csv(name, file_path, header_row):
        table = TableData(name, file_path)
        table.header_row = header_row

        raw = None
        for encoding in ['utf-8-sig', 'utf-8', 'gbk', 'gb2312', 'latin1']:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    reader = csv.reader(f)
                    raw = [row for row in reader]
                break
            except UnicodeDecodeError:
                continue
        if raw is None:
            raise ValueError(f"无法读取CSV文件: {file_path}")

        table.raw_data = raw
        _pad_rows(table.raw_data)

        header_idx = header_row - 1
        if header_idx < len(table.raw_data):
            table.headers = [str(h).strip() if h else '' for h in table.raw_data[header_idx]]
        else:
            table.headers = []

        for i, h in enumerate(table.headers):
            if not h:
                table.headers[i] = f"Col_{index_to_col_letter(i)}"

        table.build_index()
        return table

    @staticmethod
    def _load_excel(name, file_path, sheet_name, header_row):
        try:
            return TableLoader._load_excel_openpyxl(name, file_path, sheet_name, header_row)
        except Exception:
            return TableLoader._load_excel_xml(name, file_path, sheet_name, header_row)

    def _load_excel_openpyxl(name, file_path, sheet_name, header_row):
        wb = load_workbook(file_path, data_only=True, read_only=True)

        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active

        table = TableData(name, file_path)
        table.header_row = header_row

        table.raw_data = []
        for row in ws.iter_rows(values_only=True):
            table.raw_data.append([str(cell) if cell is not None else '' for cell in row])

        wb.close()
        _pad_rows(table.raw_data)

        header_idx = header_row - 1
        if header_idx < len(table.raw_data):
            table.headers = [str(h).strip() if h else '' for h in table.raw_data[header_idx]]
        else:
            table.headers = []

        for i, h in enumerate(table.headers):
            if not h:
                table.headers[i] = f"Col_{index_to_col_letter(i)}"

        table.build_index()
        return table

    def _load_excel_xml(name, file_path, sheet_name, header_row):
        import zipfile
        import xml.etree.ElementTree as ET
        import re

        table = TableData(name, file_path)
        table.header_row = header_row
        table.raw_data = []

        with zipfile.ZipFile(file_path, 'r') as z:
            if sheet_name:
                sheet_names = TableLoader._get_sheet_names_xml(file_path)
                if sheet_name in sheet_names:
                    sheet_idx = sheet_names.index(sheet_name) + 1
                else:
                    sheet_idx = 1
            else:
                sheet_idx = 1

            sheet_file = f'xl/worksheets/sheet{sheet_idx}.xml'
            if sheet_file not in z.namelist():
                for n in z.namelist():
                    if n.startswith('xl/worksheets/') and n.endswith('.xml'):
                        sheet_file = n
                        break

            shared_strings = []
            ss_path = 'xl/sharedStrings.xml'
            if ss_path in z.namelist():
                with z.open(ss_path) as f:
                    tree = ET.parse(f)
                    root = tree.getroot()
                    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
                    for si in root.iter(f'{{{ns}}}si'):
                        text_parts = []
                        for t in si.iter(f'{{{ns}}}t'):
                            if t.text:
                                text_parts.append(t.text)
                        shared_strings.append(''.join(text_parts))

            with z.open(sheet_file) as f:
                tree = ET.parse(f)
                root = tree.getroot()
                ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'

                rows_data = {}
                for row_elem in root.iter(f'{{{ns}}}row'):
                    row_num = int(row_elem.get('r', '0'))
                    cells = {}
                    for cell in row_elem.iter(f'{{{ns}}}c'):
                        ref = cell.get('r', '')
                        col_match = re.match(r'([A-Z]+)', ref)
                        if not col_match:
                            continue
                        col_letter = col_match.group(1)
                        col_idx = col_letter_to_index(col_letter)
                        cell_type = cell.get('t', '')
                        val_elem = cell.find(f'{{{ns}}}v')
                        val = val_elem.text if val_elem is not None else ''
                        if cell_type == 's' and val:
                            idx = int(val)
                            val = shared_strings[idx] if idx < len(shared_strings) else val
                        cells[col_idx] = val
                    rows_data[row_num] = cells

                if rows_data:
                    max_row = max(rows_data.keys())
                    max_col = 0
                    for cells in rows_data.values():
                        if cells:
                            max_col = max(max_col, max(cells.keys()))
                    for r in range(1, max_row + 1):
                        row = [''] * (max_col + 1)
                        cells = rows_data.get(r, {})
                        for c, v in cells.items():
                            if c <= max_col:
                                row[c] = str(v)
                        table.raw_data.append(row)

        _pad_rows(table.raw_data)

        header_idx = header_row - 1
        if header_idx < len(table.raw_data):
            table.headers = [str(h).strip() if h else '' for h in table.raw_data[header_idx]]
        else:
            table.headers = []

        for i, h in enumerate(table.headers):
            if not h:
                table.headers[i] = f"Col_{index_to_col_letter(i)}"

        table.build_index()
        return table

    @staticmethod
    def get_sheet_names(file_path):
        ext = os.path.splitext(file_path)[1].lower()
        if ext in ('.xlsx', '.xlsm'):
            try:
                wb = load_workbook(file_path, data_only=True, read_only=True)
                names = wb.sheetnames
                wb.close()
                return names
            except Exception:
                return TableLoader._get_sheet_names_xml(file_path)
        return []

    def _get_sheet_names_xml(file_path):
        import zipfile
        import xml.etree.ElementTree as ET
        names = []
        try:
            with zipfile.ZipFile(file_path, 'r') as z:
                with z.open('xl/workbook.xml') as f:
                    tree = ET.parse(f)
                    root = tree.getroot()
                    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    for sheet in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet'):
                        name = sheet.get('name', '')
                        if name:
                            names.append(name)
        except Exception:
            pass
        return names


def _pad_rows(raw_data):
    if not raw_data:
        return
    max_cols = max(len(row) for row in raw_data)
    for row in raw_data:
        while len(row) < max_cols:
            row.append('')
